"""Synthetic PPTX template behavior; needs python-pptx and pydantic (Open WebUI image).

python3 Michael/tests/test_slide_template.py
No corporate template or office documents are required by these tests.
"""
import asyncio
import base64
import importlib.util
from io import BytesIO
import json
from pathlib import Path
import sys
import unittest
import zipfile
import tempfile
from types import SimpleNamespace
from unittest.mock import patch, Mock, AsyncMock

from pptx import Presentation
from pptx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / 'tools'))
spec = importlib.util.spec_from_file_location('template_slides', ROOT / 'tools/generate_slides.py')
slides = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = slides
spec.loader.exec_module(slides)


def starter():
    prs = Presentation()
    prs.slide_layouts[0].name = 'Title Slide_White'
    prs.slide_layouts[1].name = 'Title and Content'
    prs.slide_layouts[2].name = 'Section Header_White'
    prs.slide_layouts[3].name = 'Two Column Slide'
    prs.slide_layouts[5].name = 'Title Only'
    prs.slide_layouts[6].name = 'Blank Slide'
    prs.slide_layouts[8].name = 'Title w/Image'
    # Turn the content placeholder of another layout into a genuine chart placeholder.
    chart_layout = prs.slide_layouts[7]
    chart_layout.name = 'Chart Slide'
    chart_layout.placeholders[1]._element.find('.//' + qn('p:ph')).set('type', 'chart')
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = 'SOURCE INSTRUCTIONS MUST NOT SHIP'
    slide.notes_slide.notes_text_frame.text = 'PRIVATE STARTER NOTES'
    prs.core_properties.author = 'Original template author'
    out = BytesIO(); prs.save(out)
    return out.getvalue()


class TemplateTests(unittest.TestCase):
    def setUp(self):
        self.raw = starter()
        self.tool = slides.Tools()
        self.tool.valves.starter_template_b64 = base64.b64encode(self.raw).decode()

    def build(self, items):
        data, n = self.tool._build({'title': 'Synthetic deck', 'author': 'Fixture Author', 'slides': items})
        self.assertEqual(n, len(Presentation(BytesIO(data)).slides))
        return Presentation(BytesIO(data)), data

    def test_retains_layout_library_geometry_and_discards_instruction_slides(self):
        prs, data = self.build([
            {'layout': 'cover', 'title': 'New title', 'subtitle': 'Synthetic subtitle', 'author': 'Ada'},
            {'layout': 'title_bullets', 'title': 'Evidence', 'bullets': ['First', {'text': 'Nested', 'level': 1}]},
            {'layout': 'two_column_text', 'title': 'Compare', 'columns': [
                {'heading': 'Before', 'bullets': ['Old']}, {'heading': 'After', 'bullets': ['New']} ]},
        ])
        source = Presentation(BytesIO(self.raw))
        self.assertEqual([x.name for x in prs.slide_layouts], [x.name for x in source.slide_layouts])
        self.assertEqual((prs.slide_width, prs.slide_height), (source.slide_width, source.slide_height))
        self.assertEqual(len(prs.slides), 3)
        self.assertEqual(prs.slides[0].slide_layout.name, 'Title Slide_White')
        self.assertEqual(prs.slides[0].placeholders[1].text, 'Synthetic subtitle | Ada')
        self.assertEqual(prs.slides[1].placeholders[1].text_frame.paragraphs[1].level, 1)
        self.assertIsNone(prs.slides[1].placeholders[1].text_frame.paragraphs[0].runs[0].font.name)
        self.assertEqual(prs.slides[2].placeholders[1].text, 'Before\nOld')
        self.assertEqual(prs.core_properties.author, 'Fixture Author')
        with zipfile.ZipFile(BytesIO(data)) as z:
            all_xml = b''.join(z.read(n) for n in z.namelist() if n.endswith('.xml'))
            self.assertNotIn(b'SOURCE INSTRUCTIONS MUST NOT SHIP', all_xml)
            self.assertNotIn(b'PRIVATE STARTER NOTES', all_xml)

    def test_column_headings_are_distinct_unbulleted_template_paragraphs(self):
        # Add a synthetic three-column layout; no corporate template is needed.
        from copy import deepcopy
        from pptx.util import Inches
        template = Presentation(BytesIO(self.raw))
        layout = template.slide_layouts[9]
        layout.name = 'Three Column Slide'
        layout.shapes._spTree.getparent().replace(layout.shapes._spTree,
                                deepcopy(template.slide_layouts[3].shapes._spTree))
        layout.__dict__.pop('shapes', None)
        third = deepcopy(layout.placeholders[2]._element)
        third.find('.//' + qn('p:cNvPr')).set('id', '99')
        third.find('.//' + qn('p:ph')).set('idx', '3')
        layout.shapes._spTree.insert_element_before(third, 'p:extLst')
        for i, ph in enumerate(list(layout.placeholders)[1:]):
            ph.left = Inches(0.5 + i * 3)
            ph.width = Inches(2.75)
        raw = BytesIO(); template.save(raw)
        self.tool.valves.starter_template_b64 = base64.b64encode(raw.getvalue()).decode()
        for count, alias in ((2, 'two_column_text'), (3, 'three_column_text')):
            with self.subTest(columns=count):
                columns = [{'heading': f'Section {i}',
                            'bullets': ['First point', {'text': 'Nested point', 'level': 1}]}
                           for i in range(count)]
                prs, _ = self.build([{'layout': alias, 'title': 'Sections', 'columns': columns}])
                bodies = [ph for ph in prs.slides[0].placeholders
                          if slides._placeholder_kind(ph) not in ('title', 'center_title', 'slide_number', 'date', 'footer')]
                self.assertEqual(len(bodies), count)
                for i, ph in enumerate(bodies):
                    heading, first, nested = ph.text_frame.paragraphs
                    self.assertEqual(heading.text, f'Section {i}')
                    self.assertIsNotNone(heading._p.find('./' + qn('a:pPr') + '/' + qn('a:buNone')))
                    self.assertTrue(heading.font.bold)
                    self.assertEqual(heading._p.get_or_add_pPr().get('marL'), '0')
                    self.assertEqual(heading._p.get_or_add_pPr().get('indent'), '0')
                    self.assertEqual(first.text, 'First point')
                    self.assertEqual(nested.level, 1)
                    self.assertIsNone(first._p.find('.//' + qn('a:buNone')))
                    self.assertIsNone(first.font.bold)
                    # Font family, size and colour still inherit the real template.
                    self.assertIsNone(heading.runs[0].font.name)
                    self.assertIsNone(heading.runs[0].font.size)
                    self.assertIsNone(first.runs[0].font.name)
        prs, _ = self.build([{'layout': 'two_column_text', 'columns': [
            {'bullets': ['Plain point']}, {'heading': '', 'bullets': ['Other point']}]}])
        for ph in list(prs.slides[0].placeholders)[1:]:
            paragraph = ph.text_frame.paragraphs[0]
            self.assertIsNone(paragraph.font.bold)
            self.assertIsNone(paragraph._p.find('.//' + qn('a:buNone')))

    def test_catalog_and_explicit_placeholder_text(self):
        catalog = json.loads(asyncio.run(self.tool.get_slide_layouts()))
        self.assertEqual(catalog['mode'], 'template')
        self.assertEqual(len(catalog['layouts']), 11)
        prs, _ = self.build([{'template_layout': 'Title and Content', 'placeholders': {'0': 'Exact title', '1': ['One', 'Two']}}])
        self.assertEqual(prs.slides[0].shapes.title.text, 'Exact title')
        self.assertEqual(prs.slides[0].placeholders[1].text, 'One\nTwo')

    def test_typed_bullets_and_blank_lines_do_not_double_or_empty_the_template_bullets(self):
        body = 'Intro sentence.\n\nKey points:\n\u2022 first\n- second\n* third\n\n\u2013 fourth\n5% growth\n-5 is negative'
        prs, _ = self.build([{'template_layout': 'Title and Content', 'title': 'T', 'body': body},
                             {'template_layout': 'Title and Content', 'title': 'U', 'bullets': ['\u2022 a', {'text': '- b', 'level': 1}, '  ']}])
        self.assertEqual(prs.slides[0].placeholders[1].text,
                         'Intro sentence.\nKey points:\nfirst\nsecond\nthird\nfourth\n5% growth\n-5 is negative')
        self.assertEqual(prs.slides[1].placeholders[1].text, 'a\nb')
        self.assertEqual(prs.slides[1].placeholders[1].text_frame.paragraphs[1].level, 1)
        empty, _ = self.build([{'template_layout': 'Title and Content', 'title': 'T', 'body': '\n\u2022\n'}])
        self.assertEqual(empty.slides[0].placeholders[1].text, '')

    def test_house_style_fields_the_template_cannot_draw_are_reported_not_silently_dropped(self):
        spec = {'title': 'Deck', 'theme': 'lenovo', 'slides': [
            {'layout': 'cover', 'title': 'Cover', 'eyebrow': 'PART I', 'chips': ['a']},
            {'template_layout': 'Title and Content', 'title': 'Fine', 'body': 'x'}]}
        data, _ = self.tool._build(spec)
        notes = spec['_layout_warnings']
        self.assertTrue(any('theme' in n for n in notes), notes)
        self.assertTrue(any(n.startswith('Slide 1:') and 'eyebrow' in n and 'chips' in n for n in notes), notes)
        self.assertFalse(any(n.startswith('Slide 2:') for n in notes), notes)

    def test_native_chart_and_table_are_editable(self):
        prs, _ = self.build([
            {'template_layout': 'Chart Slide', 'title': 'Data', 'labels': ['A', 'B'], 'values': [2, 3], 'chart_type': 'bar'},
            {'layout': 'table', 'title': 'Table', 'headers': ['Metric', 'Value'], 'rows': [['A', 2], ['B', 3]]},
        ])
        chart = next(s.chart for s in prs.slides[0].shapes if s.has_chart)
        self.assertEqual(list(chart.series[0].values), [2.0, 3.0])
        table = next(s.table for s in prs.slides[1].shapes if s.has_table)
        self.assertEqual(table.cell(2, 1).text, '3')

    def test_picture_is_inserted_into_native_placeholder(self):
        from PIL import Image
        image = BytesIO()
        Image.new('RGB', (40, 30), 'blue').save(image, 'PNG')
        prs, _ = self.build([{'layout': 'text_image_right', 'title': 'Synthetic image', '_img': image.getvalue()}])
        pictures = [shape for shape in prs.slides[0].shapes if hasattr(shape, 'image')]
        self.assertEqual(len(pictures), 1)
        self.assertTrue(pictures[0].is_placeholder)
        self.assertEqual(pictures[0].image.blob, image.getvalue())

    def test_attachment_id_generation_keeps_native_picture_and_aspect_ratio(self):
        from PIL import Image
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'upload.png'
            Image.new('RGBA', (120, 40), (30, 50, 200, 128)).save(path)
            files = Mock()
            files.get_file_by_id = AsyncMock(return_value=SimpleNamespace(user_id='owner', path='stored-object'))
            storage = Mock()
            storage.get_file.return_value = str(path)
            saved = []
            async def save(data, **kwargs):
                saved.append(data)
                return 'deck.pptx', '/download', None, 'result-id'
            with patch.object(slides, '_ImageFiles', files), patch.object(slides, '_ImageStorage', storage), patch.object(self.tool, '_save', save):
                result = asyncio.run(self.tool.generate_slides(json.dumps({'slides': [
                    {'layout': 'text_image_right', 'title': 'Attachment', 'image_file_id': 'upload-id'},
                    {'layout': 'text_image_right', 'image': {'file_id': 'upload-id'}},
                ]}), __user__={'id': 'owner'}))
            self.assertIn('/download', result)
            prs = Presentation(BytesIO(saved[0]))
            for slide in prs.slides:
                picture = next(s for s in slide.shapes if hasattr(s, 'image'))
                self.assertTrue(picture.is_placeholder)
                self.assertEqual(picture.image.size, (120, 40))
                with Image.open(BytesIO(picture.image.blob)) as image:
                    self.assertEqual(image.getpixel((0, 0))[3], 128)
            storage.get_file.assert_called_with('stored-object')

    def test_attachment_ownership_is_required_before_storage_access(self):
        files, storage = Mock(), Mock()
        with patch.object(slides, '_ImageFiles', files), patch.object(slides, '_ImageStorage', storage):
            for row, user in [(None, {'id': 'owner'}), (SimpleNamespace(user_id='other', path='/private'), {'id': 'owner', 'role': 'admin'}), (None, None)]:
                files.get_file_by_id.return_value = row
                with self.assertRaisesRegex(ValueError, 'unavailable or not owned'):
                    asyncio.run(slides._attachment_image('upload-id', user))
            storage.get_file.assert_not_called()

    def test_attachment_invalid_and_oversized_images_fail_without_fallback(self):
        with patch.object(slides, '_attachment_image', AsyncMock(return_value=b'not an image')):
            with self.assertRaisesRegex(ValueError, 'supported raster'):
                asyncio.run(slides._resolve_one_image({'file_id': 'id', 'url': 'https://example.com/fallback.png'}, self.tool.valves, None, {'id': 'owner'}, ratio=None))
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'too-big'; path.write_bytes(b'12345')
            files, storage = Mock(), Mock()
            files.get_file_by_id.return_value = SimpleNamespace(user_id='owner', path='stored')
            storage.get_file.return_value = str(path)
            with patch.object(slides, '_ImageFiles', files), patch.object(slides, '_ImageStorage', storage), patch.object(slides, '_MAX_IMAGE_BYTES', 4):
                with self.assertRaisesRegex(ValueError, 'byte limit'):
                    asyncio.run(slides._attachment_image('id', {'id': 'owner'}))

    def test_terminal_image_reads_and_path_restrictions(self):
        import contextlib
        import io
        import os
        with tempfile.TemporaryDirectory() as directory:
            home = Path(directory); (home / 'workspace/assets').mkdir(parents=True)
            original = b'x' * (110 * 1024)
            (home / 'workspace/assets/image.png').write_bytes(original)
            async def local_call(context, payload):
                code = workspace_delivery._TERMINAL_FILE_PROGRAM.replace('PAYLOAD', repr(base64.b64encode(json.dumps(payload).encode()).decode()))
                code = code.replace("os.path.expanduser('~')", repr(str(home)))
                out = io.StringIO()
                with contextlib.redirect_stdout(out):exec(compile(code, '<terminal fixture>', 'exec'), {})
                return json.loads(out.getvalue())
            import workspace_delivery
            with patch.object(workspace_delivery, '_terminal_file_call', local_call):
                read = asyncio.run(slides._terminal_image(None, 'assets/image.png'))
                self.assertEqual(read, original)
                for invalid in ('/etc/passwd', '../secret', 'assets/../../secret', 'assets\\secret'):
                    with self.assertRaises(ValueError):asyncio.run(slides._terminal_image(None, invalid))
                (home / 'workspace/assets/link.png').symlink_to(home / 'workspace/assets/image.png')
                with self.assertRaises(OSError):asyncio.run(slides._terminal_image(None, 'assets/link.png'))
                (home / 'workspace/linked').symlink_to(home / 'workspace/assets', target_is_directory=True)
                with self.assertRaises(OSError):asyncio.run(slides._terminal_image(None, 'linked/image.png'))
                os.mkfifo(home / 'workspace/assets/pipe')
                with self.assertRaisesRegex(ValueError, 'regular'):asyncio.run(slides._terminal_image(None, 'assets/pipe'))

    def test_terminal_image_generation_and_partial_delivery_failure(self):
        from PIL import Image
        image = BytesIO();Image.new('RGB', (120, 40), 'blue').save(image, 'PNG')
        spec = {'terminal_output': True, 'slides': [{'layout': 'text_image_right', 'title': 'From workspace', 'terminal_image_path': 'assets/example.png'}]}
        save = AsyncMock(return_value=('deck.pptx', '/download', None, 'id'))
        with patch.object(slides, '_terminal_context', AsyncMock(return_value=('url', {}, {}))) as context, patch.object(slides, '_terminal_image', AsyncMock(return_value=image.getvalue())), patch.object(slides, '_terminal_save', AsyncMock(return_value='~/workspace/output/example.pptx')) as upload, patch.object(self.tool, '_save', save):
            result = asyncio.run(self.tool.generate_slides(json.dumps(spec), __user__={'id': 'owner'}, __metadata__={'terminal_id': 'selected'}))
            self.assertIn('~/workspace/output/example.pptx', result)
            deck = Presentation(BytesIO(upload.call_args.args[1]))
            self.assertTrue(any(hasattr(s, 'image') for s in deck.slides[0].shapes))
            upload.side_effect = ValueError('failed')
            result = asyncio.run(self.tool.generate_slides(json.dumps(spec), __user__={'id': 'owner'}, __metadata__={'terminal_id': 'selected'}))
            self.assertIn('/download', result)
            self.assertIn('upload failed', result)
            context.side_effect = ValueError('access denied')
            upload.reset_mock();save.reset_mock()
            result = asyncio.run(self.tool.generate_slides(json.dumps(spec)))
            self.assertIn('access denied', result)
            upload.assert_not_called();save.assert_not_called()

    def test_table_pagination_preserves_cells_repeats_headers_and_stays_bounded(self):
        rows = [[f'Item {i}', 'Detailed synthetic explanation ' * 12, str(i)] for i in range(18)]
        prs, _ = self.build([{'layout': 'table', 'title': 'Long table', 'headers': ['Item', 'Details', 'Count'], 'rows': rows}])
        self.assertGreater(len(prs.slides), 1)
        recovered = []
        for slide in prs.slides:
            shape = next(s for s in slide.shapes if s.has_table)
            self.assertLess(shape.top + shape.height, prs.slide_height)
            self.assertEqual(shape.table.cell(0, 0).text, 'Item')
            for row in list(shape.table.rows)[1:]:
                recovered.append([' '.join(c.text.split()) for c in row.cells])
        self.assertEqual(recovered, [[str(c).strip() for c in r] for r in rows])

    def test_wide_table_and_extra_columns_continue_without_losing_content(self):
        prs, _ = self.build([{'layout': 'table', 'headers': ['Key']+[f'C{i}' for i in range(8)], 'rows': [['a']+list(range(8))]},
                             {'layout': 'two_column_text', 'columns': ['One', 'Two', 'Three', 'Four']}])
        self.assertEqual(len(prs.slides), 4)
        self.assertEqual(prs.slides[1].shapes[-1].table.cell(1, 0).text, 'a')
        self.assertIn('Four', '\n'.join(s.text for s in prs.slides[3].shapes if s.has_text_frame))

    def test_pie_labels_and_many_category_fallback(self):
        from pptx.enum.chart import XL_CHART_TYPE
        prs, _ = self.build([{'layout': 'chart', 'chart_type': 'pie', 'labels': ['A', 'B', 'C'], 'values': [2, 3, 4]},
                             {'layout': 'chart', 'chart_type': 'pie', 'labels': [str(i) for i in range(9)], 'values': list(range(1,10))}])
        charts = [next(s.chart for s in slide.shapes if s.has_chart) for slide in prs.slides]
        self.assertTrue(charts[0].has_legend)
        self.assertTrue(charts[0].plots[0].data_labels.show_category_name)
        self.assertTrue(charts[0].plots[0].data_labels.show_percentage)
        self.assertEqual(charts[1].chart_type, XL_CHART_TYPE.BAR_CLUSTERED)
        self.assertEqual(list(charts[1].series[0].values), list(range(1,10)))

    def test_very_tall_row_and_chart_caption_geometry(self):
        text = '\n'.join(f'Line {i}' for i in range(70))
        prs, _ = self.build([{'layout': 'table', 'headers': ['Details'], 'rows': [[text]]},
                             {'layout': 'chart', 'labels': ['A','B'], 'values': [1,2], 'insight': 'Short interpretation'}])
        recovered=[]
        for slide in list(prs.slides)[:-1]:
            shape=next(s for s in slide.shapes if s.has_table)
            self.assertLess(shape.top+shape.height,prs.slide_height)
            recovered.extend(shape.table.cell(r,0).text for r in range(1,len(shape.table.rows)))
        self.assertEqual('\n'.join(recovered),text)
        chart=next(s for s in prs.slides[-1].shapes if s.has_chart)
        self.assertGreater(chart.width,0)
        self.assertGreater(chart.height,0)

    def test_closing_takeaways_are_preserved_before_artwork(self):
        source=Presentation(BytesIO(self.raw));source.slide_layouts[6].name='Closing Slide'
        out=BytesIO();source.save(out);self.tool.valves.starter_template_b64=base64.b64encode(out.getvalue()).decode()
        prs,_=self.build([{'layout':'closing','title':'Next steps','takeaways':['Review evidence','Confirm owners'],'contact':'Example team'}])
        self.assertEqual(len(prs.slides),2)
        self.assertEqual(prs.slides[-1].slide_layout.name,'Closing Slide')
        content='\n'.join(s.text for s in prs.slides[0].shapes if s.has_text_frame)
        for text in ['Next steps','Review evidence','Confirm owners','Example team']:self.assertIn(text,content)

    def test_json_comments_and_trailing_commas_preserve_strings(self):
        raw = '{"url":"https://example.test/a//b",/* note */"values":[1,2,],// note\n}'
        self.assertEqual(slides._json_spec(raw), {'url':'https://example.test/a//b','values':[1,2]})
        with self.assertRaises(ValueError):slides._json_spec('{"values":[1,unknown]}')

    def test_bootstrap_template_loading_and_invalid_package(self):
        import tempfile
        sys.path.insert(0, str(ROOT / 'bootstrap'))
        import office_tools
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'starter.pptx'
            self.assertIsNone(office_tools.load_starter_template(path))
            path.write_bytes(self.raw)
            self.assertEqual(office_tools.load_starter_template(path), self.raw)
            path.write_bytes(b'not a presentation')
            with self.assertRaises(ValueError):
                office_tools.load_starter_template(path)

    def test_bad_specs_fail_instead_of_silently_using_another_design(self):
        bad = [
            {'template_layout': 'Does not exist'},
            {'layout': 'kpi_row', 'stats': [{'value': 2}]},
            {'layout': 'title_body', 'placeholders': {'999': 'Unknown'}},
            {'layout': 'title_bullets', 'bullets': [{'text': 'Bad', 'level': 8}]},
            {'layout': 'blank', 'title': 'Would be silently lost'},
            {'template_layout': 'Chart Slide', 'labels': ['A', 'B'], 'values': [1]},
            {'template_layout': 'Chart Slide', 'labels': ['A'], 'values': ['unknown']},
            {'template_layout': 'Chart Slide', 'labels': ['A'], 'values': [float('nan')]},
        ]
        for item in bad:
            with self.subTest(item=item), self.assertRaises(ValueError):
                self.build([item])

    def test_invalid_template_and_embedded_objects_are_refused(self):
        self.tool.valves.starter_template_b64 = 'not-base64'
        with self.assertRaises(ValueError):
            self.build([{'layout': 'cover'}])
        out = BytesIO(self.raw)
        with zipfile.ZipFile(out, 'a') as z:
            z.writestr('ppt/embeddings/oleObject1.bin', b'synthetic')
        self.tool.valves.starter_template_b64 = base64.b64encode(out.getvalue()).decode()
        with self.assertRaisesRegex(ValueError, 'embedded objects'):
            self.build([{'layout': 'cover'}])

    def test_no_template_keeps_legacy_renderer_available(self):
        self.tool.valves.starter_template_b64 = ''
        self.assertEqual(json.loads(asyncio.run(self.tool.get_slide_layouts()))['mode'], 'legacy')
        prs, _ = self.build([{'layout': 'cover', 'title': 'Legacy'}])
        self.assertEqual(len(prs.slides), 1)



def table_shapes(slide):
    return [shape for shape in slide.shapes if getattr(shape, 'has_table', False) and shape.has_table]


class DeckQualityTests(unittest.TestCase):
    """What the deck in the failing conversation looked like, and what must now happen instead."""
    MONTHS = ['2026-07', '2026-08', '2026-09', '2026-10']

    def setUp(self):
        self.tool = slides.Tools()
        self.tool.valves.starter_template_b64 = base64.b64encode(starter()).decode()

    def build(self, items):
        spec = {'title': 'Deck', 'slides': items}
        report = spec.setdefault('_layout_warnings', [])
        slides._promote_text_tables(items, report)
        data, _ = self.tool._build(spec)
        return Presentation(BytesIO(data)), data, spec['_layout_warnings']

    @staticmethod
    def png(width, height, color='navy'):
        from PIL import Image
        out = BytesIO(); Image.new('RGB', (width, height), color).save(out, 'PNG'); return out.getvalue()

    # -- tables written as text --------------------------------------------------------------

    def test_pipe_separated_text_becomes_a_real_table(self):
        rows = 'Series | Jul | Aug\nM Series | 1 | 9\nChromebooks | 1 | 0'
        prs, _, report = self.build([{'template_layout': 'Title and Content', 'title': 'Counts', 'body': rows}])
        (table,) = table_shapes(prs.slides[0])
        self.assertEqual([c.text for c in table.table.rows[0].cells], ['Series', 'Jul', 'Aug'])
        self.assertEqual(table.table.cell(2, 0).text, 'Chromebooks')
        self.assertFalse(any('|' in shape.text_frame.text for shape in prs.slides[0].shapes if shape.has_text_frame))
        self.assertTrue(any('real table' in line for line in report))

    def test_markdown_tables_and_bulleted_rows_are_promoted_too(self):
        prs, _, _ = self.build([{'layout': 'title_body', 'title': 'MD', 'body': '| A | B |\n|---|:--:|\n| 1 | 2 |'}])
        self.assertEqual([c.text for c in table_shapes(prs.slides[0])[0].table.rows[0].cells], ['A', 'B'])
        prs, _, _ = self.build([{'template_layout': 'Title and Content', 'title': 'B', 'bullets': ['A | B', '1 | 2']}])
        self.assertEqual(len(table_shapes(prs.slides[0])), 1)

    def test_ordinary_text_with_a_pipe_is_left_alone(self):
        for body in ('Either this | or that', 'one | two\nthree', 'a | b\nc | d | e', '| only\n| one column', 'Plain sentence.\nAnother.'):
            self.assertIsNone(slides._text_table(body), body)

    # -- native charts -----------------------------------------------------------------------

    def chart_spec(self, count, **extra):
        return {'template_layout': 'Chart Slide', 'title': 'Chart', 'chart_type': 'bar', 'labels': self.MONTHS,
                'datasets': [{'label': f'Series {i}', 'data': [(i + j) % 5 + 1 for j in range(4)]} for i in range(count)], **extra}

    def test_ten_series_are_stacked_with_distinct_colours(self):
        prs, _, report = self.build([self.chart_spec(10)])
        chart = next(s.chart for s in prs.slides[0].shapes if getattr(s, 'has_chart', False) and s.has_chart)
        self.assertEqual(str(chart.chart_type).split()[0], 'COLUMN_STACKED')
        colours = [str(series.format.fill.fore_color.rgb) for series in chart.series]
        self.assertEqual(len(set(colours)), 10)
        self.assertTrue(any('Stacked a chart with 10 series' in line for line in report))

    def test_a_few_series_stay_side_by_side_and_can_be_forced_either_way(self):
        chart = lambda items: next(s.chart for s in self.build(items)[0].slides[0].shapes if getattr(s, 'has_chart', False) and s.has_chart)
        self.assertEqual(str(chart([self.chart_spec(3)]).chart_type).split()[0], 'COLUMN_CLUSTERED')
        self.assertEqual(str(chart([self.chart_spec(3, stacked=True)]).chart_type).split()[0], 'COLUMN_STACKED')
        self.assertEqual(str(chart([self.chart_spec(8, stacked=False)]).chart_type).split()[0], 'COLUMN_CLUSTERED')

    def test_labels_hide_zeros_and_thin_stacked_segments_are_deleted_not_blanked(self):
        prs, data, _ = self.build([self.chart_spec(8)])
        with zipfile.ZipFile(BytesIO(data)) as z:
            xml = next(z.read(n).decode() for n in z.namelist() if n.startswith('ppt/charts/chart') and n.endswith('.xml'))
        self.assertIn('formatCode="0;-0;;"', xml)
        self.assertIn('<c:delete val="1"/>', xml)
        self.assertRegex(xml, r'<c:legendPos(?: val="r")?/>')  # right is the schema default, so it may be omitted

    # -- an exported chart image -------------------------------------------------------------

    def test_an_exported_chart_is_shown_whole_inside_the_chart_area(self):
        wide = self.png(1600, 686)
        prs, _, _ = self.build([{'template_layout': 'Chart Slide', 'title': 'Exported', '_img': wide}])
        slide = prs.slides[0]
        pictures = [s for s in slide.shapes if s.shape_type == 13]
        self.assertEqual(len(pictures), 1)
        self.assertFalse(any(getattr(s, 'is_placeholder', False) and s.placeholder_format.type is not None and 'CHART' in str(s.placeholder_format.type) for s in slide.shapes))
        picture = pictures[0]
        self.assertAlmostEqual(picture.width / picture.height, 1600 / 686, delta=0.02)
        self.assertEqual((picture.crop_left, picture.crop_right, picture.crop_top, picture.crop_bottom), (0, 0, 0, 0))
        area = [p for p in slide.slide_layout.placeholders if 'CHART' in str(p.placeholder_format.type)][0]
        self.assertGreaterEqual(picture.left, area.left); self.assertGreaterEqual(picture.top, area.top)
        self.assertLessEqual(picture.left + picture.width, area.left + area.width + 1)
        self.assertLessEqual(picture.top + picture.height, area.top + area.height + 1)

    def test_a_tall_image_is_fitted_not_cropped_or_stretched(self):
        prs, _, _ = self.build([{'template_layout': 'Chart Slide', 'title': 'Tall', '_img': self.png(600, 1200)}])
        picture = [s for s in prs.slides[0].shapes if s.shape_type == 13][0]
        self.assertAlmostEqual(picture.width / picture.height, 0.5, delta=0.01)

    def test_a_picture_layout_crops_by_default_and_can_contain(self):
        crop = lambda fit: (lambda p: (p.crop_left + p.crop_right + p.crop_top + p.crop_bottom))(
            [s for s in self.build([{'template_layout': 'Title w/Image', 'title': 'T', '_img': self.png(1600, 400), **fit}])[0].slides[0].shapes if s.shape_type in (13, 14) and hasattr(s, 'crop_left')][0])
        self.assertGreater(crop({}), 0)
        self.assertEqual(crop({'image_fit': 'contain'}), 0)

    def test_an_image_on_a_layout_with_no_place_for_it_is_still_refused(self):
        with self.assertRaisesRegex(ValueError, 'picture placeholder'):
            self.build([{'template_layout': 'Title and Content', 'title': 'No picture area', 'image_url': 'https://example.com/x.png'}])

    # -- tables are sized by content ---------------------------------------------------------

    def test_a_table_with_a_wide_text_column_stays_on_one_slide(self):
        names = ['M Series desktops (ThinkCentre)', 'Lenovo Chromebooks Series', 'L Series laptops (ThinkPad)', '500 Series laptops (ideapad)',
                 'P Series workstations (ThinkStation)', 'T Series laptops (ThinkPad)', 'X Series laptops (ThinkPad)', 'P Series laptops (ThinkPad)',
                 '300 Series laptops (ideapad)', 'Edge Series laptops (ThinkPad)']
        rows = [[n, 1, 9, 19, 1, 30] for n in names]
        prs, _, _ = self.build([{'template_layout': 'Title and Content', 'title': 'Counts', 'headers': ['Product Series', 'Jul', 'Aug', 'Sep', 'Oct', 'Total'], 'rows': rows}])
        self.assertEqual(len(prs.slides), 1)
        table = table_shapes(prs.slides[0])[0].table
        widths = [c.width for c in table.columns]
        self.assertGreater(widths[0], 3 * widths[1])
        self.assertEqual(len(table.rows), 11)

    def test_a_long_table_still_continues_onto_more_slides(self):
        rows = [[f'Item {i}', 'Detailed synthetic explanation ' * 8, str(i)] for i in range(40)]
        prs, _, report = self.build([{'template_layout': 'Title and Content', 'title': 'Long', 'headers': ['Item', 'Details', 'Count'], 'rows': rows}])
        self.assertGreater(len(prs.slides), 1)
        self.assertTrue(any('continued across' in line for line in report))

    # -- the file itself ---------------------------------------------------------------------

    def test_the_package_keeps_default_namespaces_so_strict_readers_can_open_it(self):
        _, data, _ = self.build([{'layout': 'cover', 'title': 'T', 'subtitle': 'S'}])
        with zipfile.ZipFile(BytesIO(data)) as z:
            for name in ('[Content_Types].xml', '_rels/.rels'):
                xml = z.read(name).decode()
                self.assertNotIn('ns0:', xml, name)
                self.assertRegex(xml, r'<(Types|Relationships) xmlns="http', name)
            self.assertFalse([n for n in z.namelist() if n.startswith('docProps/thumbnail')])
            self.assertEqual(z.namelist()[0], '[Content_Types].xml')

if __name__ == '__main__':
    unittest.main()
