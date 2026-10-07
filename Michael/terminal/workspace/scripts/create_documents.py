"""Build simple Word/PPTX source packets from a directory without overwriting outputs."""
import argparse
from pathlib import Path
from docx import Document
from docx.shared import Inches as DocInches
from pptx import Presentation
from pptx.util import Inches, Pt
from PIL import Image

def create(source, output):
    source = Path(source).expanduser().resolve()
    output = Path(output).expanduser().resolve()
    if not source.is_dir():
        raise ValueError('Source must be a directory')
    files = sorted(p for p in source.iterdir() if p.is_file() and not p.is_symlink()
                   and p.suffix.lower() in {'.txt', '.md', '.png', '.jpg', '.jpeg'})
    if not files or len(files) > 30:
        raise ValueError('Provide 1–30 supported top-level source files')
    texts = {p: p.read_text(encoding='utf-8') for p in files if p.suffix.lower() in {'.txt', '.md'}}
    if any(len(t) > 20000 for t in texts.values()):
        raise ValueError('Text exceeds 20,000 characters; prepare a deliberate summary first')
    # Refuse existing delivery directories to preserve prior artifacts.
    output.mkdir(parents=True, exist_ok=False)
    doc = Document(); doc.add_heading('Source packet', 0)
    deck = Presentation(); deck.slide_width = Inches(13.333); deck.slide_height = Inches(7.5)
    title = deck.slides.add_slide(deck.slide_layouts[0]); title.shapes.title.text = 'Source packet'
    title.placeholders[1].text = 'Prepared from supplied files; no generated factual claims'
    for p in files:
        doc.add_heading(p.name, level=1)
        if p in texts:
            text = texts[p]; doc.add_paragraph(text)
            chunks = [text[i:i+1000] for i in range(0,len(text),1000)] or ['(empty file)']
            for n, chunk in enumerate(chunks, 1):
                slide = deck.slides.add_slide(deck.slide_layouts[1])
                slide.shapes.title.text = f'{p.name} ({n}/{len(chunks)})'
                frame = slide.placeholders[1].text_frame; frame.text = chunk
                for paragraph in frame.paragraphs: paragraph.font.size = Pt(18)
        else:
            doc.add_picture(str(p), width=DocInches(5.5))
            slide = deck.slides.add_slide(deck.slide_layouts[5]); slide.shapes.title.text = p.name
            with Image.open(p) as im: w,h=im.size
            scale=min(11/w,5.4/h); width,height=w*scale,h*scale
            slide.shapes.add_picture(str(p), Inches((13.333-width)/2), Inches(1.6), width=Inches(width), height=Inches(height))
    doc.save(output/'source-packet.docx'); deck.save(output/'source-packet.pptx')
    (output/'sources.txt').write_text('\n'.join(p.name for p in files)+'\n')
    print('Created Word and PowerPoint source packets in', output)

if __name__ == '__main__':
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--source',required=True);p.add_argument('--output',required=True)
    a=p.parse_args();create(a.source,a.output)
